"""ワイン会 実施前後レポート（Excel）を生成するスクリプト。

使い方:
    pip install openpyxl
    python3 wine-event-report/build_report.py

出力: wine-event-report/ワイン会レポート.xlsx
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("ワイン会レポート.xlsx")

# ブランドカラー（instagram-operation-manual 5-1）
BORDEAUX = "5A1425"
GOLD = "D4A24E"
WHITE = "FAF6F2"
DARK = "2A1418"
INPUT_FILL = "FFF6D9"   # 入力セル（淡い黄）
CALC_FILL = "F2F2F2"    # 自動計算セル（淡いグレー）
FONT = "Meiryo"

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def f(bold=False, color=DARK, size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)


F_BASE = f()
F_INPUT = f(color="0000FF")
F_HEAD = f(bold=True, color="FFFFFF")
F_SECTION = f(bold=True, color=BORDEAUX, size=12)
F_TITLE = f(bold=True, color=BORDEAUX, size=16)
F_NOTE = f(color="7F7F7F", size=9)

FILL_HEAD = PatternFill("solid", fgColor=BORDEAUX)
FILL_INPUT = PatternFill("solid", fgColor=INPUT_FILL)
FILL_CALC = PatternFill("solid", fgColor=CALC_FILL)

YEN = '#,##0;-#,##0;"-"'
PCT = '0.0%;-0.0%;"-"'
NUM = '#,##0;-#,##0;"-"'
DEC = '0.0;-0.0;"-"'
DATE = 'yyyy/mm/dd'
TIME = 'hh:mm'

WIDTHS = {  # イベントシートの列幅（A〜S）
    "A": 5, "B": 22, "C": 14, "D": 14, "E": 12, "F": 16, "G": 12, "H": 11,
    "I": 16, "J": 12, "K": 12, "L": 11, "M": 12, "N": 10, "O": 12,
    "P": 11, "Q": 10, "R": 12, "S": 28,
}

CHANNELS = ["Instagram", "Meta広告", "LINE", "知人・紹介", "以前参加",
            "イベントサイト", "メールマガジン", "検索", "その他"]
STATUSES = ["申込", "参加", "キャンセル", "無断欠席"]


def style(c, font=F_BASE, fill=None, fmt=None, align=None, border=True):
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align or Alignment(vertical="center", wrap_text=True)
    if border:
        c.border = BORDER


def inp(ws, ref, value=None, fmt=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    style(c, F_INPUT, FILL_INPUT, fmt)
    return c


def calc(ws, ref, formula, fmt=None):
    c = ws[ref]
    c.value = formula
    style(c, F_BASE, FILL_CALC, fmt)
    return c


def label(ws, ref, text, bold=False):
    c = ws[ref]
    c.value = text
    style(c, f(bold=bold))
    return c


def head_row(ws, row, headers, start_col=1):
    for i, h in enumerate(headers):
        c = ws.cell(row=row, column=start_col + i, value=h)
        style(c, F_HEAD, FILL_HEAD,
              align=Alignment(horizontal="center", vertical="center", wrap_text=True))


def section(ws, row, text, note=None):
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_SECTION
    ws.cell(row=row, column=1).border = Border(bottom=Side(style="medium", color=GOLD))
    for col in range(2, 20):
        ws.cell(row=row, column=col).border = Border(bottom=Side(style="medium", color=GOLD))
    if note:
        n = ws.cell(row=row, column=6, value=note)
        n.font = F_NOTE


def merge_input(ws, rng, value=None, fmt=None, height=None):
    ws.merge_cells(rng)
    first = rng.split(":")[0]
    c = inp(ws, first, value, fmt)
    c.alignment = Alignment(vertical="top", wrap_text=True)
    for row in ws[rng]:
        for cell in row:
            cell.border = BORDER
            cell.fill = FILL_INPUT
    if height:
        ws.row_dimensions[ws[first].row].height = height
    return c


# ---------------------------------------------------------------------------
# イベントシート（1開催 = 1シート）
# サマリーはこのセル番地を INDIRECT で参照するため、レイアウトを変えたら
# CELL を必ず更新すること。
# ---------------------------------------------------------------------------
CELL = {}


def build_event_sheet(ws, data=None):
    data = data or {}
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False

    ws["A1"] = "ワイン会 実施前後レポート"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("黄色セル＝入力 ／ グレーセル＝自動計算（触らない）。"
                "実施前は【計画】列、実施後は【実績】列と参加者リストを記入。")
    ws["A2"].font = F_NOTE

    # ---------------- 基本情報 ----------------
    r = 4
    section(ws, r, "1. 基本情報")
    basics = [
        ("no", "回No", None),
        ("name", "イベント名", None),
        ("type", "種別", None),
        ("date", "開催日", DATE),
        ("open", "開場", TIME),
        ("start", "開始", TIME),
        ("end", "終了", TIME),
        ("venue", "会場", None),
        ("capacity", "定員（名）", NUM),
        ("fee_pre", "会費・事前決済（円/税込）", YEN),
        ("fee_day", "会費・当日決済（円/税込）", YEN),
        ("owner", "責任者（企画・最終判断）", None),
        ("staff", "運営担当", None),
        ("status", "ステータス", None),
    ]
    r += 1
    for key, text, fmt in basics:
        label(ws, f"B{r}", text, bold=True)
        ws.merge_cells(f"C{r}:E{r}")
        inp(ws, f"C{r}", data.get(key), fmt)
        for col in "DE":
            ws[f"{col}{r}"].border = BORDER
            ws[f"{col}{r}"].fill = FILL_INPUT
        CELL[key] = f"C{r}"
        r += 1
    ws[CELL["date"]].alignment = Alignment(horizontal="left", vertical="center")
    dv_status = DataValidation(type="list", formula1='"企画中,募集中,実施済,中止"', allow_blank=True)
    dv_type = DataValidation(type="list",
                             formula1='"経営者交流（Class-E）,for Beginner,その他"', allow_blank=True)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_type)
    dv_status.add(CELL["status"])
    dv_type.add(CELL["type"])

    # 実施前チェック（基本情報の右側）
    cr = 5
    ws.merge_cells(f"G4:K4")
    ws["G4"] = "実施前チェックリスト"
    ws["G4"].font = f(bold=True, color=BORDEAUX)
    head_row(ws, cr, ["項目", "", "", "期限", "状況"], start_col=7)
    ws.merge_cells(f"G{cr}:I{cr}")
    checklist = data.get("checklist", [
        "会場予約の確定",
        "提供ワイン確定・発注",
        "告知ページ／LP公開",
        "Instagram告知投稿",
        "LINE告知（3週間前）",
        "LINE告知（1週間前）",
        "LINE告知（前日）",
        "広告入稿・審査通過",
        "参加者へ前日リマインド",
        "受付名簿・領収書の準備",
        "写真掲載の同意確認方法を準備",
        "当日の流れ・台本の確定",
    ])
    dv_check = DataValidation(type="list", formula1='"未着手,対応中,完了,対象外"', allow_blank=True)
    ws.add_data_validation(dv_check)
    for i, item in enumerate(checklist):
        rr = cr + 1 + i
        ws.merge_cells(f"G{rr}:I{rr}")
        inp(ws, f"G{rr}", item)
        for col in "HI":
            ws[f"{col}{rr}"].border = BORDER
            ws[f"{col}{rr}"].fill = FILL_INPUT
        inp(ws, f"J{rr}", None, DATE)
        inp(ws, f"K{rr}", "未着手")
        dv_check.add(f"K{rr}")
    CELL["check_first"] = cr + 1
    CELL["check_last"] = cr + len(checklist)
    ws.merge_cells(f"M4:O4")
    ws["M4"] = "準備の進捗"
    ws["M4"].font = f(bold=True, color=BORDEAUX)
    label(ws, "M5", "完了", bold=True)
    calc(ws, "N5",
         f'=COUNTIF(K{CELL["check_first"]}:K{CELL["check_last"]},"完了")'
         f'&" / "&(COUNTA(K{CELL["check_first"]}:K{CELL["check_last"]})'
         f'-COUNTIF(K{CELL["check_first"]}:K{CELL["check_last"]},"対象外"))')

    # 以降のセクションの行番号を先に決める（参加者リストの範囲が必要なため）
    r_kpi = r + 1                     # 2. 計画vs実績
    kpi_rows = 14
    r_channel = r_kpi + 2 + kpi_rows + 1   # 3. 集客チャネル
    r_money = r_channel + 2 + len(CHANNELS) + 1 + 1  # 4. 収支
    money_items = 16
    r_ads = r_money + 2 + money_items + 1  # 5. 広告
    ad_rows = 4
    r_wine = r_ads + 2 + ad_rows + 1 + 1   # 6. 提供ワイン
    wine_rows = 10
    r_time = r_wine + 2 + wine_rows + 1 + 1  # 7. タイムテーブル
    time_rows = 8
    r_review = r_time + 2 + time_rows + 1  # 8. 振り返り
    review_items = 6
    r_part = r_review + 1 + review_items + 1  # 9. 参加者リスト
    part_rows = 40
    p1 = r_part + 2
    p2 = p1 + part_rows - 1

    P = {  # 参加者リストの列
        "name": "B", "kana": "C", "sex": "D", "age": "E", "company": "F",
        "title": "G", "repeat": "H", "channel": "I", "applied": "J",
        "pay": "K", "status": "L", "fee": "M", "paid": "N", "sales": "O",
        "photo": "P", "sat": "Q", "next": "R", "note": "S",
    }

    def rng(key):
        col = P[key]
        return f"${col}${p1}:${col}${p2}"

    joined = f'{rng("status")},"参加"'

    # ---------------- 2. 計画 vs 実績 ----------------
    section(ws, r_kpi, "2. 計画 vs 実績（KPI）",
            "実施前に【計画】、実施後に【実績】。実績の多くは参加者リストから自動集計")
    h = r_kpi + 1
    head_row(ws, h, ["", "項目", "計画（実施前）", "実績（実施後）", "差異", "達成率", "メモ"])
    ws.merge_cells(f"G{h}:K{h}")
    plan = data.get("plan", {})
    kpis = [
        ("apply", "申込数（名）", NUM, f'=COUNTA({rng("name")})'),
        ("join", "参加者数（名）", NUM, f'=COUNTIF({joined})'),
        ("male", "　うち男性", NUM, f'=COUNTIFS({joined},{rng("sex")},"男性")'),
        ("female", "　うち女性", NUM, f'=COUNTIFS({joined},{rng("sex")},"女性")'),
        ("new", "　うち新規", NUM, f'=COUNTIFS({joined},{rng("repeat")},"新規")'),
        ("repeat", "　うちリピーター", NUM, f'=COUNTIFS({joined},{rng("repeat")},"リピーター")'),
        ("cancel", "キャンセル数", NUM, f'=COUNTIF({rng("status")},"キャンセル")'),
        ("noshow", "無断欠席数", NUM, f'=COUNTIF({rng("status")},"無断欠席")'),
        ("fill", "充足率（参加者÷定員）", PCT, None),
        ("line", "LINE友だち追加数（告知期間）", NUM, "INPUT"),
        ("sat", "満足度 平均（5段階）", DEC,
         f'=IFERROR(AVERAGEIFS({rng("sat")},{joined},{rng("sat")},">0"),0)'),
        ("next", "次回参加意向「あり」率", PCT,
         f'=IFERROR(COUNTIFS({joined},{rng("next")},"あり")/COUNTIF({joined}),0)'),
        ("photo", "写真掲載 同意数", NUM, f'=COUNTIFS({joined},{rng("photo")},"同意")'),
        ("profit_pp", "参加者1名あたり粗利（円）", YEN, None),
    ]
    rr = h + 1
    for key, text, fmt, actual in kpis:
        CELL["kpi_" + key] = rr
        label(ws, f"B{rr}", text, bold=not text.startswith("　"))
        if key == "fill":
            calc(ws, f"C{rr}", f'=IFERROR(C{CELL["kpi_join"]}/{CELL["capacity"]},0)', fmt)
            calc(ws, f"D{rr}", f'=IFERROR(D{CELL["kpi_join"]}/{CELL["capacity"]},0)', fmt)
        elif key == "profit_pp":
            pass  # 収支セクション確定後に設定
        else:
            inp(ws, f"C{rr}", plan.get(key), fmt)
            if actual == "INPUT":
                inp(ws, f"D{rr}", None, fmt)
            else:
                calc(ws, f"D{rr}", actual, fmt)
        calc(ws, f"E{rr}", f'=IF(C{rr}="","",D{rr}-C{rr})', fmt)
        if fmt in (PCT, DEC):
            calc(ws, f"F{rr}", '=""')
        else:
            calc(ws, f"F{rr}", f'=IF(OR(C{rr}="",C{rr}=0),"",D{rr}/C{rr})', PCT)
        ws.merge_cells(f"G{rr}:K{rr}")
        inp(ws, f"G{rr}", data.get("kpi_memo", {}).get(key))
        for col in "HIJK":
            ws[f"{col}{rr}"].border = BORDER
            ws[f"{col}{rr}"].fill = FILL_INPUT
        rr += 1

    # ---------------- 3. 集客チャネル ----------------
    section(ws, r_channel, "3. 集客チャネル別（知ったきっかけ）",
            "受付で全員に「どこで知りましたか」を聞く（運用マニュアル9-3）")
    h = r_channel + 1
    head_row(ws, h, ["", "チャネル", "計画（名）", "申込（名）", "参加（名）",
                     "男性（参加）", "女性（参加）", "構成比（参加）", "メモ"])
    ws.merge_cells(f"I{h}:K{h}")
    ch_plan = data.get("channel_plan", {})
    first_ch = h + 1
    last_ch = h + len(CHANNELS)
    for i, ch in enumerate(CHANNELS):
        rr = h + 1 + i
        label(ws, f"B{rr}", ch)
        plan_v, memo = ch_plan.get(ch, (None, None))
        inp(ws, f"C{rr}", plan_v, NUM)
        calc(ws, f"D{rr}", f'=COUNTIF({rng("channel")},B{rr})', NUM)
        calc(ws, f"E{rr}", f'=COUNTIFS({joined},{rng("channel")},B{rr})', NUM)
        calc(ws, f"F{rr}", f'=COUNTIFS({joined},{rng("channel")},B{rr},{rng("sex")},"男性")', NUM)
        calc(ws, f"G{rr}", f'=COUNTIFS({joined},{rng("channel")},B{rr},{rng("sex")},"女性")', NUM)
        calc(ws, f"H{rr}", f'=IFERROR(E{rr}/E${last_ch + 1},0)', PCT)
        ws.merge_cells(f"I{rr}:K{rr}")
        inp(ws, f"I{rr}", memo)
        for col in "JK":
            ws[f"{col}{rr}"].border = BORDER
            ws[f"{col}{rr}"].fill = FILL_INPUT
    tr = last_ch + 1
    label(ws, f"B{tr}", "合計", bold=True)
    for col in "CDEFG":
        calc(ws, f"{col}{tr}", f"=SUM({col}{first_ch}:{col}{last_ch})", NUM)
        ws[f"{col}{tr}"].font = f(bold=True)
    calc(ws, f"H{tr}", f"=SUM(H{first_ch}:H{last_ch})", PCT)
    calc(ws, f"I{tr}", f'=IF(D{tr}<>D{CELL["kpi_apply"]},"※きっかけ未入力の申込者あり","")')
    ws.merge_cells(f"I{tr}:K{tr}")
    CELL["ch_ad_row"] = first_ch + CHANNELS.index("Meta広告")

    # ---------------- 5. 広告（先に行番号を確定） ----------------
    ad_h = r_ads + 1
    ad_first = ad_h + 1
    ad_last = ad_h + ad_rows
    ad_total = ad_last + 1

    # ---------------- 6. 提供ワイン（行番号） ----------------
    w_h = r_wine + 1
    w_first = w_h + 1
    w_last = w_h + wine_rows
    w_total = w_last + 1

    # ---------------- 4. 収支 ----------------
    section(ws, r_money, "4. 収支（円・税込）",
            "計画は実施前の見込み。実績の会費・物販・ワイン仕入・広告費は自動集計")
    h = r_money + 1
    head_row(ws, h, ["", "項目", "計画（実施前）", "実績（実施後）", "差異", "", "メモ"])
    ws.merge_cells(f"G{h}:K{h}")
    money_plan = data.get("money_plan", {})
    items = [
        ("s_fee", "売上：会費収入", f'=SUM({rng("fee")})'),
        ("s_goods", "売上：ワイン物販", f'=SUM({rng("sales")})'),
        ("s_other", "売上：その他", "INPUT"),
        ("s_total", "売上 合計", "SUM_S"),
        ("c_wine", "費用：ワイン仕入", f"=I{w_total}"),
        ("c_food", "費用：おつまみ・フード", "INPUT"),
        ("c_venue", "費用：会場費", "INPUT"),
        ("c_ads", "費用：広告費", f"=E{ad_total}"),
        ("c_fee", "費用：決済手数料", "INPUT"),
        ("c_staff", "費用：人件費・謝礼", "INPUT"),
        ("c_goods", "費用：備品・印刷・消耗品", "INPUT"),
        ("c_other", "費用：その他", "INPUT"),
        ("c_total", "費用 合計", "SUM_C"),
        ("profit", "粗利（売上−費用）", "PROFIT"),
        ("margin", "粗利率", "MARGIN"),
        ("cpa", "広告CPA（広告費÷広告経由参加者）", "CPA"),
    ]
    rr = h + 1
    for key, text, actual in items:
        CELL["m_" + key] = rr
        rr += 1
    s_first, s_last = CELL["m_s_fee"], CELL["m_s_other"]
    c_first, c_last = CELL["m_c_wine"], CELL["m_c_other"]
    for key, text, actual in items:
        rr = CELL["m_" + key]
        is_total = key in ("s_total", "c_total", "profit")
        label(ws, f"B{rr}", text, bold=is_total or key in ("margin", "cpa"))
        fmt = PCT if key == "margin" else YEN
        if actual == "SUM_S":
            calc(ws, f"C{rr}", f"=SUM(C{s_first}:C{s_last})", fmt)
            calc(ws, f"D{rr}", f"=SUM(D{s_first}:D{s_last})", fmt)
        elif actual == "SUM_C":
            calc(ws, f"C{rr}", f"=SUM(C{c_first}:C{c_last})", fmt)
            calc(ws, f"D{rr}", f"=SUM(D{c_first}:D{c_last})", fmt)
        elif actual == "PROFIT":
            for col in "CD":
                calc(ws, f"{col}{rr}",
                     f'={col}{CELL["m_s_total"]}-{col}{CELL["m_c_total"]}', fmt)
        elif actual == "MARGIN":
            for col in "CD":
                calc(ws, f"{col}{rr}",
                     f'=IFERROR({col}{CELL["m_profit"]}/{col}{CELL["m_s_total"]},0)', fmt)
        elif actual == "CPA":
            inp(ws, f"C{rr}", money_plan.get(key), fmt)
            calc(ws, f"D{rr}",
                 f'=IFERROR(D{CELL["m_c_ads"]}/E{CELL["ch_ad_row"]},0)', fmt)
        else:
            if key == "s_fee" and "s_fee" not in money_plan:
                calc(ws, f"C{rr}",
                     f'=C{CELL["kpi_join"]}*{CELL["fee_pre"]}', fmt)
                ws[f"C{rr}"].comment = Comment(
                    "既定値＝計画参加者数×事前決済の会費。見込みが違う場合は上書きしてよい。", "report")
            else:
                inp(ws, f"C{rr}", money_plan.get(key), fmt)
            if actual == "INPUT":
                inp(ws, f"D{rr}", None, fmt)
            else:
                calc(ws, f"D{rr}", actual, fmt)
        if is_total:
            for col in "BCD":
                ws[f"{col}{rr}"].font = f(bold=True)
        if key in ("margin",):
            calc(ws, f"E{rr}", '=""')
        else:
            calc(ws, f"E{rr}", f'=D{rr}-C{rr}', fmt)
        ws[f"F{rr}"].border = BORDER
        ws.merge_cells(f"G{rr}:K{rr}")
        inp(ws, f"G{rr}", data.get("money_memo", {}).get(key))
        for col in "HIJK":
            ws[f"{col}{rr}"].border = BORDER
            ws[f"{col}{rr}"].fill = FILL_INPUT

    # 1名あたり粗利（KPI表）
    rr = CELL["kpi_profit_pp"]
    for col in "CD":
        calc(ws, f"{col}{rr}",
             f'=IFERROR({col}{CELL["m_profit"]}/{col}{CELL["kpi_join"]},0)', YEN)

    # ---------------- 5. 広告 ----------------
    section(ws, r_ads, "5. 広告実績（Meta広告など）",
            "広告を出さない回は空欄のままでよい。記録の残し方はMeta広告ガイド9-5")
    head_row(ws, ad_h, ["", "広告セット", "配信期間", "予算（円）", "消化額（円）",
                        "リード数", "CPL（円）", "広告経由参加", "CPA（円）", "メモ"])
    ws.merge_cells(f"J{ad_h}:K{ad_h}")
    ad_sets = data.get("ads", [])
    for i in range(ad_rows):
        rr = ad_first + i
        a = ad_sets[i] if i < len(ad_sets) else {}
        inp(ws, f"B{rr}", a.get("set"))
        inp(ws, f"C{rr}", a.get("period"))
        inp(ws, f"D{rr}", a.get("budget"), YEN)
        inp(ws, f"E{rr}", None, YEN)
        inp(ws, f"F{rr}", None, NUM)
        calc(ws, f"G{rr}", f"=IFERROR(E{rr}/F{rr},0)", YEN)
        inp(ws, f"H{rr}", None, NUM)
        calc(ws, f"I{rr}", f"=IFERROR(E{rr}/H{rr},0)", YEN)
        ws.merge_cells(f"J{rr}:K{rr}")
        inp(ws, f"J{rr}", a.get("memo"))
        ws[f"K{rr}"].border = BORDER
        ws[f"K{rr}"].fill = FILL_INPUT
    label(ws, f"B{ad_total}", "合計", bold=True)
    calc(ws, f"C{ad_total}", '=""')
    for col in "DEFH":
        calc(ws, f"{col}{ad_total}", f"=SUM({col}{ad_first}:{col}{ad_last})",
             NUM if col in "FH" else YEN)
    calc(ws, f"G{ad_total}", f"=IFERROR(E{ad_total}/F{ad_total},0)", YEN)
    calc(ws, f"I{ad_total}", f"=IFERROR(E{ad_total}/H{ad_total},0)", YEN)

    # ---------------- 6. 提供ワイン ----------------
    section(ws, r_wine, "6. 提供ワイン",
            "参加者の反応が良かった銘柄は次回・物販の判断材料にする")
    head_row(ws, w_h, ["No", "銘柄", "生産者", "産地", "品種", "タイプ",
                       "仕入単価（円）", "本数", "小計（円）", "反応・評価メモ"])
    ws.merge_cells(f"J{w_h}:K{w_h}")
    wines = data.get("wines", [])
    for i in range(wine_rows):
        rr = w_first + i
        wv = wines[i] if i < len(wines) else {}
        c = ws.cell(row=rr, column=1, value=i + 1)
        style(c, align=Alignment(horizontal="center"))
        inp(ws, f"B{rr}", wv.get("name"))
        inp(ws, f"C{rr}", wv.get("maker"))
        inp(ws, f"D{rr}", wv.get("region"))
        inp(ws, f"E{rr}", wv.get("grape"))
        inp(ws, f"F{rr}", wv.get("type"))
        inp(ws, f"G{rr}", None, YEN)
        inp(ws, f"H{rr}", None, NUM)
        calc(ws, f"I{rr}", f"=G{rr}*H{rr}", YEN)
        ws.merge_cells(f"J{rr}:K{rr}")
        inp(ws, f"J{rr}", wv.get("memo"))
        ws[f"K{rr}"].border = BORDER
        ws[f"K{rr}"].fill = FILL_INPUT
    label(ws, f"B{w_total}", "合計", bold=True)
    calc(ws, f"H{w_total}", f"=SUM(H{w_first}:H{w_last})", NUM)
    calc(ws, f"I{w_total}", f"=SUM(I{w_first}:I{w_last})", YEN)
    dv_wtype = DataValidation(type="list",
                              formula1='"白,赤,ロゼ,スパークリング,オレンジ,甘口・貴腐,その他"',
                              allow_blank=True)
    ws.add_data_validation(dv_wtype)
    dv_wtype.add(f"F{w_first}:F{w_last}")

    # ---------------- 7. タイムテーブル ----------------
    section(ws, r_time, "7. 当日の流れ（タイムテーブル）")
    t_h = r_time + 1
    head_row(ws, t_h, ["", "時刻", "内容", "", "", "", "担当", "実際の進行・メモ"])
    ws.merge_cells(f"C{t_h}:F{t_h}")
    ws.merge_cells(f"H{t_h}:K{t_h}")
    tl = data.get("timeline", [])
    for i in range(time_rows):
        rr = t_h + 1 + i
        tv = tl[i] if i < len(tl) else (None, None)
        inp(ws, f"B{rr}", tv[0], TIME)
        ws.merge_cells(f"C{rr}:F{rr}")
        inp(ws, f"C{rr}", tv[1])
        inp(ws, f"G{rr}")
        ws.merge_cells(f"H{rr}:K{rr}")
        inp(ws, f"H{rr}")
        for col in "DEFIJK":
            ws[f"{col}{rr}"].border = BORDER
            ws[f"{col}{rr}"].fill = FILL_INPUT

    # ---------------- 8. 振り返り ----------------
    section(ws, r_review, "8. 振り返り（実施後）",
            "開催後3日以内に記入 → 責任者が翌回の方針を判断")
    reviews = [
        "良かった点",
        "課題・トラブル",
        "参加者の声（アンケート・会話から）",
        "集客の振り返り（どの階層で詰まったか）",
        "次回の改善アクション（3つまで）",
        "責任者の判断・次回方針",
    ]
    for i, text in enumerate(reviews):
        rr = r_review + 1 + i
        label(ws, f"B{rr}", text, bold=True)
        merge_input(ws, f"C{rr}:K{rr}", height=48)

    # ---------------- 9. 参加者リスト ----------------
    section(ws, r_part, "9. 参加者リスト",
            "個人情報を含むため、記入済みファイルはGitHubに置かず、閲覧権限を絞ったGoogleドライブで管理")
    head_row(ws, r_part + 1, [
        "No", "氏名", "フリガナ", "性別", "年代", "会社・所属", "役職",
        "新規/リピート", "知ったきっかけ", "申込日", "決済方法", "ステータス",
        "会費受領額（円）", "入金確認", "物販購入額（円）", "写真掲載", "満足度（1-5）",
        "次回参加意向", "備考（アレルギー・領収書宛名など）",
    ])
    ws.row_dimensions[r_part + 1].height = 32
    for i in range(part_rows):
        rr = p1 + i
        c = ws.cell(row=rr, column=1, value=i + 1)
        style(c, align=Alignment(horizontal="center"))
        for key, col in P.items():
            fmt = {"applied": DATE, "fee": YEN, "sales": YEN, "sat": NUM}.get(key)
            inp(ws, f"{col}{rr}", None, fmt)
    lists = {
        "sex": '"男性,女性,その他・回答なし"',
        "age": '"20代,30代,40代,50代,60代以上"',
        "repeat": '"新規,リピーター"',
        "channel": '"' + ",".join(CHANNELS) + '"',
        "pay": '"クレジットカード,銀行振込,当日現金,その他"',
        "status": '"' + ",".join(STATUSES) + '"',
        "paid": '"済,未"',
        "photo": '"同意,不可"',
        "sat": '"1,2,3,4,5"',
        "next": '"あり,未定,なし"',
    }
    for key, formula in lists.items():
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{P[key]}{p1}:{P[key]}{p2}")

    ws.freeze_panes = "A4"
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = GOLD
    CELL.update({
        "p1": p1, "p2": p2,
        "sales": f'D{CELL["m_s_total"]}', "cost": f'D{CELL["m_c_total"]}',
        "profit": f'D{CELL["m_profit"]}', "ads": f'D{CELL["m_c_ads"]}',
        "plan_join": f'C{CELL["kpi_join"]}',
        "plan_sales": f'C{CELL["m_s_total"]}',
    })
    return P


# ---------------------------------------------------------------------------
# 各回のデータ（リポジトリ内の資料で確認できた「計画値」のみ。実績は空欄）
# ---------------------------------------------------------------------------
EVENT_0525 = {
    "no": "—",
    "name": "経営者交流ハンガリーワイン会（Class-E）",
    "type": "経営者交流（Class-E）",
    "date": "2026-05-25",
    "open": "18:30", "start": "19:00", "end": "21:00",
    "venue": "The Place ラウンジ（渋谷区道玄坂2-11-1 JMFビル渋谷03 5F）",
    "capacity": 20,
    "fee_pre": 15000, "fee_day": 16000,
    "owner": "高橋", "staff": None, "status": None,
    "plan": {"join": 20},
    "kpi_memo": {"join": "LPは「20名限定」表記、料金欄は「定員20〜30名」表記（class-e-lp/index.html）"},
    "channel_plan": {},
    "wines": [
        {"name": "プレミアム フルミント", "maker": "Steigler", "region": "Sopron", "type": "白"},
        {"name": "エリジウム", "maker": "Steigler", "region": "Hungary", "type": "赤"},
        {"name": "トゥラン", "maker": "カローイムレ", "region": "Hungary", "type": "赤"},
        {"name": "プレステージ ブリュット", "maker": "KREINBACHER", "region": "Somló",
         "type": "スパークリング"},
        {"name": "ハールシュレベリュー", "maker": "Steigler", "region": "Hungary", "type": "白"},
        {"name": "トカイアスー 6プットニョシュ 2018", "maker": "Gotz", "region": "Tokaj",
         "type": "甘口・貴腐"},
        {"name": "ライナイリーズリング", "maker": "タクレル", "region": "Hungary", "type": "白"},
    ],
    "timeline": [
        ("18:30", "開場・ウェルカムドリンク"),
        ("18:45", "主催挨拶・ハンガリーワイン解説"),
        ("19:00", "テイスティング & 食事・自由対話"),
        ("20:30", "閉会"),
        ("21:00", "完全撤収"),
    ],
}

EVENT_0831 = {
    "no": "—",
    "name": "ハンガリーワイン会 for Beginner",
    "type": "for Beginner",
    "date": "2026-08-31",
    "open": None, "start": "19:00", "end": None,
    "venue": "六本木一丁目「経営者の部室」",
    "capacity": 30,
    "fee_pre": 7000, "fee_day": None,
    "owner": "高橋", "staff": "秘書", "status": None,
    "plan": {"join": 30, "male": 20, "female": 10},
    "kpi_memo": {
        "join": "集客目標30名（Meta広告ガイド冒頭）",
        "male": "男女比は内部目標。告知・広告には出さない（運用マニュアル4-1）",
    },
    "channel_plan": {
        "Meta広告": (13, "計画11〜15名の中央値（Meta広告ガイド5-3）"),
        "知人・紹介": (17, "ツテ・オーガニック計画15〜19名の中央値（同5-3）"),
    },
    "money_plan": {"c_ads": 30000, "cpa": 3000},
    "money_memo": {
        "c_ads": "総予算30,000円（Meta広告ガイド5-1）",
        "cpa": "基準3,000円／許容上限5,000円（同5-4）",
    },
    "ads": [
        {"set": "男性セット", "period": "8/7〜8/30", "budget": 9000},
        {"set": "女性セット", "period": "8/7〜8/30", "budget": 17000},
        {"set": "リターゲティング", "period": "8/24〜8/30", "budget": 3000},
    ],
}

SAMPLE = {
    "no": 99,
    "name": "（記入例）ハンガリーワイン会 for Beginner",
    "type": "for Beginner",
    "date": "2026-10-26",
    "open": "18:45", "start": "19:00", "end": "21:00",
    "venue": "六本木一丁目「経営者の部室」",
    "capacity": 30, "fee_pre": 7000, "fee_day": 7500,
    "owner": "高橋", "staff": "秘書", "status": "実施済",
    "plan": {"apply": 32, "join": 28, "male": 18, "female": 10, "new": 16, "repeat": 12,
             "line": 25, "sat": 4.5, "next": 0.6, "photo": 20},
    "channel_plan": {"Instagram": (8, None), "LINE": (6, None), "知人・紹介": (10, None),
                     "以前参加": (4, None)},
    "money_plan": {"s_goods": 20000, "c_wine": 45000, "c_food": 30000, "c_venue": 30000,
                   "c_fee": 3000, "c_ads": 0},
    "wines": [{"name": "（例）フルミント", "maker": "（例）○○ワイナリー", "region": "Tokaj",
               "grape": "フルミント", "type": "白", "memo": "一番人気。物販で3本売れた"}],
    "timeline": [("18:45", "開場・受付（知ったきっかけを確認）"), ("19:00", "乾杯・ハンガリーワイン紹介"),
                 ("20:45", "次回案内・LINE登録のお願い"), ("21:00", "終了")],
}
SAMPLE_PEOPLE = [
    ["山田 太郎", "ヤマダ タロウ", "男性", "30代", "株式会社サンプル", "営業部長", "新規",
     "Instagram", "2026-10-10", "クレジットカード", "参加", 7000, "済", 5000, "同意", 5, "あり",
     "領収書宛名：株式会社サンプル"],
    ["佐藤 花子", "サトウ ハナコ", "女性", "20代", "", "", "リピーター", "LINE", "2026-10-15",
     "当日現金", "参加", 7500, "済", 0, "不可", 4, "未定", "甲殻類アレルギー"],
    ["鈴木 一郎", "スズキ イチロウ", "男性", "40代", "", "", "新規", "知人・紹介", "2026-10-20",
     "銀行振込", "キャンセル", 0, "未", 0, "", None, "", "前日キャンセル"],
]


def fill_values(ws, data):
    """テンプレート上の入力セルへ値を入れる（文字列の日付・時刻を型変換）。"""
    import datetime as dt
    for key in ("date",):
        if data.get(key):
            ws[CELL[key]].value = dt.date.fromisoformat(data[key])
    for key in ("open", "start", "end"):
        if data.get(key):
            h, m = data[key].split(":")
            ws[CELL[key]].value = dt.time(int(h), int(m))
    for row in ws.iter_rows(min_col=2, max_col=2):
        c = row[0]
        if isinstance(c.value, str) and ":" in c.value and len(c.value) == 5 and c.value[:2].isdigit():
            h, m = c.value.split(":")
            c.value = dt.time(int(h), int(m))


# ---------------------------------------------------------------------------
# サマリーシート
# ---------------------------------------------------------------------------
def build_summary(ws, sheet_names):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = BORDEAUX
    ws["A1"] = "ワイン会 開催サマリー"
    ws["A1"].font = F_TITLE
    ws["A2"] = ("B列に各回のシート名を入れると、そのシートから自動で集計されます。"
                "新しい回は「テンプレート」シートをコピー → シート名を変更 → ここに1行追加。")
    ws["A2"].font = F_NOTE

    # KPIカード
    cards = [
        ("開催回数（実施済）", '=COUNTIF($X$8:$X$37,"実施済")', NUM),
        ("累計参加者（名）", "=SUM($K$8:$K$37)", NUM),
        ("累計売上（円）", "=SUM($Q$8:$Q$37)", YEN),
        ("累計粗利（円）", "=SUM($S$8:$S$37)", YEN),
        ("平均参加者／回", '=IFERROR(SUMIF($X$8:$X$37,"実施済",$K$8:$K$37)/COUNTIF($X$8:$X$37,"実施済"),0)', DEC),
        ("平均満足度", '=IFERROR(AVERAGEIF($W$8:$W$37,">0"),0)', DEC),
    ]
    col = 2
    for text, formula, fmt in cards:
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 2)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 2)
        c = ws.cell(row=4, column=col, value=text)
        style(c, F_HEAD, FILL_HEAD, align=Alignment(horizontal="center", vertical="center"))
        v = ws.cell(row=5, column=col, value=formula)
        style(v, f(bold=True, color=BORDEAUX, size=16), PatternFill("solid", fgColor=WHITE), fmt,
              Alignment(horizontal="center", vertical="center"))
        for cc in range(col, col + 3):
            ws.cell(row=4, column=cc).border = BORDER
            ws.cell(row=5, column=cc).border = BORDER
        col += 3
    ws.row_dimensions[5].height = 34

    headers = [
        "回No", "シート名（入力）", "開催日", "曜日", "イベント名", "種別", "会場", "定員",
        "申込数", "参加者 計画", "参加者 実績", "男性", "女性", "新規", "リピーター",
        "充足率", "売上（円）", "費用（円）", "粗利（円）", "粗利率", "1名あたり粗利（円）",
        "広告費（円）", "満足度", "ステータス",
    ]
    head_row(ws, 7, headers)
    ws.row_dimensions[7].height = 32

    def ref(cell):
        return f'INDIRECT("\'"&$B{{r}}&"\'!{cell}")'

    cols = [
        ("A", CELL["no"], None),
        ("C", CELL["date"], DATE),
        ("E", CELL["name"], None),
        ("F", CELL["type"], None),
        ("G", CELL["venue"], None),
        ("H", CELL["capacity"], NUM),
        ("I", f'D{CELL["kpi_apply"]}', NUM),
        ("J", CELL["plan_join"], NUM),
        ("K", f'D{CELL["kpi_join"]}', NUM),
        ("L", f'D{CELL["kpi_male"]}', NUM),
        ("M", f'D{CELL["kpi_female"]}', NUM),
        ("N", f'D{CELL["kpi_new"]}', NUM),
        ("O", f'D{CELL["kpi_repeat"]}', NUM),
        ("Q", CELL["sales"], YEN),
        ("R", CELL["cost"], YEN),
        ("S", CELL["profit"], YEN),
        ("V", CELL["ads"], YEN),
        ("W", f'D{CELL["kpi_sat"]}', DEC),
        ("X", CELL["status"], None),
    ]
    first, last = 8, 37
    for r in range(first, last + 1):
        name = sheet_names[r - first] if r - first < len(sheet_names) else None
        inp(ws, f"B{r}", name)
        for c_letter, cell, fmt in cols:
            formula = ref(cell).format(r=r)
            if fmt is None:
                v = f'=IF($B{r}="","",IFERROR({formula}&"",""))'
            else:
                v = f'=IF($B{r}="","",IFERROR({formula},""))'
            calc(ws, f"{c_letter}{r}", v, fmt)
        ws[f"C{r}"].alignment = Alignment(horizontal="center", vertical="center")
        calc(ws, f"D{r}",
             f'=IF(OR($B{r}="",C{r}=""),"",CHOOSE(WEEKDAY(C{r}),"日","月","火","水","木","金","土"))')
        ws[f"D{r}"].alignment = Alignment(horizontal="center", vertical="center")
        calc(ws, f"P{r}", f'=IF(OR($B{r}="",N(H{r})=0),"",K{r}/H{r})', PCT)
        calc(ws, f"T{r}", f'=IF(OR($B{r}="",N(Q{r})=0),"",S{r}/Q{r})', PCT)
        calc(ws, f"U{r}", f'=IF(OR($B{r}="",N(K{r})=0),"",S{r}/K{r})', YEN)
    tr = last + 1
    label(ws, f"B{tr}", "合計・平均", bold=True)
    for c_letter in "HIJKLMNOQRSV":
        fmt = YEN if c_letter in "QRSV" else NUM
        calc(ws, f"{c_letter}{tr}", f"=SUM({c_letter}{first}:{c_letter}{last})", fmt)
        ws[f"{c_letter}{tr}"].font = f(bold=True)
    calc(ws, f"P{tr}", f"=IFERROR(K{tr}/H{tr},0)", PCT)
    calc(ws, f"T{tr}", f"=IFERROR(S{tr}/Q{tr},0)", PCT)
    calc(ws, f"U{tr}", f"=IFERROR(S{tr}/K{tr},0)", YEN)
    calc(ws, f"W{tr}", f'=IFERROR(AVERAGEIF(W{first}:W{last},">0"),0)', DEC)

    widths = {"A": 6, "B": 20, "C": 11, "D": 5, "E": 30, "F": 16, "G": 26, "H": 7,
              "I": 8, "J": 9, "K": 9, "L": 7, "M": 7, "N": 7, "O": 9, "P": 8,
              "Q": 12, "R": 12, "S": 12, "T": 8, "U": 12, "V": 11, "W": 8, "X": 9}
    for k, w in widths.items():
        ws.column_dimensions[k].width = w
    ws.freeze_panes = "C8"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws[f"B{tr + 2}"] = ("※ 売上・費用・参加者などの実績は各回シートの【実績】列から自動取得。"
                        "未実施の回は計画値のみ表示されます。")
    ws[f"B{tr + 2}"].font = F_NOTE


def build_guide(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 90
    ws["B1"] = "使い方"
    ws["B1"].font = F_TITLE
    rows = [
        ("シート構成", "サマリー：全開催の一覧と累計 ／ 各回シート：1開催＝1シート ／ "
                   "テンプレート：新しい回の雛形 ／ 記入例：書き方の見本（集計対象外）"),
        ("セルの色", "黄色＝入力セル（青文字）／ グレー＝自動計算（数式。上書きしない）"),
        ("新しい回を追加する",
         "①「テンプレート」シートを右クリック→コピー ②シート名を「YYYYMMDD_種別」に変更 "
         "③サマリーB列の空き行にそのシート名を入力（自動で集計される）"),
        ("実施前（告知開始〜前日）",
         "1.基本情報／チェックリスト、2.KPIの【計画】列、3.チャネル別の計画、4.収支の【計画】列、"
         "5.広告の予算、6.提供ワイン、7.タイムテーブルを記入。参加者リストには申込が入り次第追加"
         "（ステータス＝申込）"),
        ("当日",
         "受付でステータスを「参加／無断欠席」に更新し、会費受領額・知ったきっかけ・写真掲載の同意を記入"),
        ("実施後（3日以内）",
         "満足度・次回参加意向・物販額、収支の【実績】（会場費など手入力分）、広告の消化額・リード数、"
         "8.振り返りを記入 → 責任者（高橋）が次回方針を判断"),
        ("自動集計の定義",
         "申込数＝参加者リストの氏名入力数 ／ 参加者数＝ステータス「参加」／ 会費収入＝会費受領額の合計 ／ "
         "ワイン仕入＝6.提供ワインの小計合計 ／ 広告費＝5.広告の消化額合計 ／ "
         "広告CPA＝広告費÷「知ったきっかけ＝Meta広告」の参加者数"),
        ("男女比の扱い",
         "男女別の目標・実績は内部管理用。告知・広告・SNSには出さない（運用マニュアル4-1）"),
        ("個人情報の管理",
         "参加者リストに個人情報を記入したファイルはGitHubに置かない。"
         "閲覧権限を責任者・運営担当に限定したGoogleドライブで管理する（運用マニュアル2章⑥・3-1）"),
        ("Googleスプレッドシートで使う場合",
         "Googleドライブにアップロード →「Googleスプレッドシートで開く」→「ファイル」→"
         "「Googleスプレッドシートとして保存」。INDIRECT関数でサマリーが連動します"),
        ("既存2回分の入力状況",
         "20260525_ClassE・20260831_Beginner は、リポジトリ内の資料（LP・Meta広告ガイド）で確認できた"
         "計画値のみ入力済み。実績（参加者・売上・費用）は資料に無いため空欄。"
         "この2回以外の開催履歴は資料に無く、未作成"),
    ]
    r = 3
    for k, v in rows:
        a = ws.cell(row=r, column=2, value=k)
        style(a, f(bold=True, color="FFFFFF"), FILL_HEAD)
        b = ws.cell(row=r, column=3, value=v)
        style(b)
        ws.row_dimensions[r].height = 46
        r += 1
    r += 1
    ws.cell(row=r, column=2, value="凡例").font = f(bold=True, color=BORDEAUX)
    style(ws.cell(row=r + 1, column=2, value="入力セル"), F_INPUT, FILL_INPUT)
    style(ws.cell(row=r + 2, column=2, value="自動計算セル"), F_BASE, FILL_CALC)


def main():
    wb = Workbook()
    guide = wb.active
    guide.title = "使い方"
    summary = wb.create_sheet("サマリー")

    event_sheets = [("20260525_ClassE", EVENT_0525), ("20260831_Beginner", EVENT_0831)]
    P = None
    for name, data in event_sheets:
        ws = wb.create_sheet(name)
        P = build_event_sheet(ws, data)
        fill_values(ws, data)
        ws["A1"] = f"ワイン会 実施前後レポート｜{data['name']}"

    tpl = wb.create_sheet("テンプレート")
    build_event_sheet(tpl, {})
    tpl.sheet_properties.tabColor = "BFBFBF"

    ex = wb.create_sheet("記入例")
    build_event_sheet(ex, SAMPLE)
    fill_values(ex, SAMPLE)
    ex["A1"] = "【記入例】ワイン会 実施前後レポート（架空のデータ・集計対象外）"
    ex.sheet_properties.tabColor = "BFBFBF"
    import datetime as dt
    for i, person in enumerate(SAMPLE_PEOPLE):
        rr = CELL["p1"] + i
        for key, value in zip(P.keys(), person):
            if key == "applied" and value:
                value = dt.date.fromisoformat(value)
            ex[f"{P[key]}{rr}"].value = value
    for i, k in enumerate(["完了"] * 10 + ["対象外", "完了"]):
        ex[f"K{CELL['check_first'] + i}"].value = k
    ex[f"D{CELL['kpi_line']}"].value = 22

    build_summary(summary, [n for n, _ in event_sheets])
    build_guide(guide)
    wb.active = 1
    wb.save(OUT)
    print(f"saved: {OUT}")


if __name__ == "__main__":
    main()
